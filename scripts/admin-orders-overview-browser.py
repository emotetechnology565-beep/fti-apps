"""Isolated admin overview acceptance: no live backend calls or order mutations."""
import base64
import json
import threading
import time
from datetime import datetime, timedelta
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

root = Path.cwd()
directory = root / 'artifacts/admin-orders-preview'
out = root / 'artifacts/admin-orders'
out.mkdir(exist_ok=True)

class Handler(SimpleHTTPRequestHandler):
    def do_GET(self):
        if not Path(self.translate_path(self.path.split('?')[0])).is_file() and '.' not in self.path.split('?')[0].rsplit('/', 1)[-1]:
            self.path = '/index.html'
        super().do_GET()
    def log_message(self, *args):
        pass

server = ThreadingHTTPServer(('127.0.0.1', 0), partial(Handler, directory=str(directory)))
threading.Thread(target=server.serve_forever, daemon=True).start()
base = f'http://127.0.0.1:{server.server_port}'
actor = '00000000-0000-4000-8000-000000000001'
enc = lambda value: base64.urlsafe_b64encode(json.dumps(value).encode()).decode().rstrip('=')
user = {'id': actor, 'aud': 'authenticated', 'role': 'authenticated', 'phone': '919000000001', 'app_metadata': {}, 'user_metadata': {}}
session = {'user': user, 'access_token': enc({'alg': 'HS256'}) + '.' + enc({'sub': actor, 'aud': 'authenticated', 'role': 'authenticated', 'aal': 'aal2', 'exp': int(time.time()) + 3600}) + '.isolated', 'refresh_token': 'isolated', 'expires_at': int(time.time()) + 3600, 'expires_in': 3600, 'token_type': 'bearer'}
identity = {'profile': {'id': actor, 'full_name': 'Preview staff', 'phone': '+919000000001', 'account_status': 'active'}, 'roles': ['super_admin'], 'permissions': ['staff.access', 'order.manage', 'operations.manage'], 'businesses': [], 'deletion_requested': False}
rows = []
for index in range(31):
    kind = ['material', 'fastener', 'commerce'][index % 3]
    status = ['pending_confirmation', 'confirmed', 'processing', 'shipped', 'delivered', 'cancelled', 'rejected'][index % 7]
    created = datetime(2026, 10, 30, 10) - timedelta(days=index % 8)
    rows.append({'id': f'10000000-0000-4000-8000-{index+1:012}', 'kind': kind, 'reference': {'material': 'RM', 'fastener': 'FN', 'commerce': 'O'}[kind] + f'-{index+1:06}', 'status': status, 'created_at': created.isoformat() + 'Z', 'buyer_name': 'Sample buyer', 'phone': '+919876543210', 'business': 'Sample Engineering', 'location': 'Industrial Road, Ludhiana, Punjab', 'material': 'Bright steel bar' if kind == 'material' else 'Hex bolt M12' if kind == 'fastener' else 'Tractor parts', 'quantity': '500 kg' if kind == 'material' else '40 pieces', 'grade': 'EN8', 'size': '40 mm', 'total': '47000', 'amount_due': None, 'closure_note': None, 'fulfillment': {'status': status, 'revision': 2, 'confirmed_total': '47000', 'confirmation_note': 'Availability and delivery terms agreed.', 'history': [], 'carrier': None}})
requests, errors = [], []
mode = {'fail': False, 'overflow': False}

def api(route):
    req = route.request
    body = (req.post_data_json or {}) if req.method == 'POST' else {}
    direct = '/rest/v1/rpc/' in req.url
    op = req.url.rsplit('/', 1)[-1] if direct else body.get('operation')
    args = body if direct else body.get('arguments', {})
    result = []
    if '/auth/v1/user' in req.url:
        result = user
    elif op == 'fti_identity_me':
        result = identity
    elif op == 'fti_order_payments':
        row = next(r for r in rows if r['id'] == args['p_id'])
        result = {'total': row['total'] if row['status'] != 'pending_confirmation' else None, 'paid': '0', 'remaining': row['total'] if row['status'] != 'pending_confirmation' else None, 'status': row['status'], 'receipts': []}
    elif op == 'fti_admin_orders':
        requests.append(args)
        if mode['fail']:
            return route.fulfill(status=500, content_type='application/json', body=json.dumps({'message': 'Preview service unavailable'}), headers={'Access-Control-Allow-Origin': '*'})
        data = [r for r in rows if args.get('p_kind', 'all') in ['all', r['kind']] and args.get('p_status', 'all') in ['all', r['status']]]
        search = args.get('p_search', '').lower()
        data = [r for r in data if not search or search in json.dumps(r).lower()]
        data = [r for r in data if (not args.get('p_from') or r['created_at'][:10] >= args['p_from']) and (not args.get('p_to') or r['created_at'][:10] <= args['p_to'])]
        summary = {'total': len(data), 'active': sum(r['status'] not in ['delivered', 'cancelled', 'rejected'] for r in data), 'awaiting': sum(r['status'] in ['pending_confirmation', 'pending_payment'] for r in data), 'processing': sum(r['status'] == 'processing' for r in data), 'shipped': sum(r['status'] == 'shipped' for r in data), 'delivered': sum(r['status'] == 'delivered' for r in data), 'cancelled': sum(r['status'] == 'cancelled' for r in data), 'rejected': sum(r['status'] == 'rejected' for r in data), 'confirmed_value': str(sum(47000 for r in data if r['status'] in ['confirmed', 'processing', 'shipped', 'delivered']))}
        if mode['overflow']:
            summary['total'] = 10001
        offset = args.get('p_offset', 0)
        result = {'orders': data[offset:offset + args.get('p_limit', 24)], 'summary': summary}
    if '/functions/v1/' in req.url:
        result = {'data': result}
    route.fulfill(status=200, content_type='application/json', body=json.dumps(result), headers={'Access-Control-Allow-Origin': '*', 'Access-Control-Allow-Headers': '*'})

try:
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        for width in [320, 390, 1440]:
            context = browser.new_context(viewport={'width': width, 'height': 1000}, accept_downloads=True)
            context.route('https://*.supabase.co/**', api)
            context.add_init_script('localStorage.setItem("fti-auth-v2",' + json.dumps(json.dumps(session)) + ')')
            page = context.new_page()
            page.on('pageerror', lambda error: errors.append(str(error)))
            page.set_default_timeout(15000)
            page.goto(base + '/admin?section=orders-overview', wait_until='networkidle')
            expect(page.get_by_role('button', name='View order ', exact=False)).to_have_count(24)
            expect(page.get_by_text('Matching orders', exact=True)).to_be_visible()
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            page.screenshot(path=str(out / f'overview-{width}.png'))
            page.get_by_role('button', name='Next orders', exact=True).click()
            expect(page.get_by_role('button', name='View order ', exact=False)).to_have_count(7)
            assert requests[-1]['p_offset'] == 24
            page.get_by_label('Search orders', exact=True).fill('RM-000001')
            page.get_by_role('button', name='Search & apply filters', exact=True).click()
            expect(page.get_by_role('button', name='View order RM-000001', exact=True)).to_be_visible()
            expect(page.get_by_role('button', name='View order ', exact=False)).to_have_count(1)
            assert requests[-1]['p_offset'] == 0
            page.get_by_role('button', name='View order RM-000001', exact=True).click()
            expect(page.get_by_text('Buyer & delivery', exact=True)).to_be_visible()
            expect(page.get_by_text('Order fulfillment', exact=True)).to_be_visible()
            page.get_by_role('button', name='Reset filters', exact=True).click()
            page.get_by_label('From date (YYYY-MM-DD)', exact=True).fill('2026-10-30')
            page.get_by_label('To date (YYYY-MM-DD)', exact=True).fill('2026-10-29')
            page.get_by_role('button', name='Search & apply filters', exact=True).click()
            expect(page.get_by_text('From date must be on or before to date.', exact=True)).to_be_visible()
            page.get_by_label('To date (YYYY-MM-DD)', exact=True).fill('2026-10-30')
            page.get_by_role('button', name='Search & apply filters', exact=True).click()
            expect(page.get_by_role('button', name='View order ', exact=False)).to_have_count(4)
            assert requests[-1]['p_from'] == '2026-10-30' and requests[-1]['p_to'] == '2026-10-30'
            page.get_by_role('button', name='Order type: Fasteners', exact=True).click()
            page.get_by_role('button', name='Order status: Confirmed', exact=True).click()
            page.get_by_role('button', name='Search & apply filters', exact=True).click()
            expect(page.get_by_text('No matching orders', exact=True)).to_be_visible()
            expect(page.get_by_role('button', name='Export Excel', exact=True)).to_be_disabled()
            page.get_by_role('button', name='Reset filters', exact=True).click()
            expect(page.get_by_role('button', name='View order ', exact=False)).to_have_count(24)
            with page.expect_download() as download:
                page.get_by_role('button', name='Export Excel', exact=True).click()
            path = out / f'orders-{width}.xlsx'
            download.value.save_as(str(path))
            workbook = load_workbook(path)
            assert workbook.sheetnames == ['Orders', 'Overview']
            assert workbook['Orders'].max_row == 32, 'Export must include more than the 24 visible rows'
            assert requests[-1]['p_limit'] == 10000 and requests[-1]['p_offset'] == 0
            mode['overflow'] = True
            page.get_by_role('button', name='Export Excel', exact=True).click()
            expect(page.get_by_text('More than 10,000 orders match. Narrow the date range or filters before exporting.', exact=True)).to_be_visible()
            mode['overflow'] = False
            mode['fail'] = True
            page.get_by_role('button', name='Refresh orders', exact=True).click()
            expect(page.get_by_text('Orders unavailable', exact=True)).to_be_visible()
            mode['fail'] = False
            page.get_by_role('button', name='Retry orders', exact=True).click()
            expect(page.get_by_role('button', name='View order ', exact=False)).to_have_count(24)
            context.close()
        browser.close()
    assert not errors, errors
    print('PASS: admin overview search, dates, type/status filters, pagination, expandable details, full Excel download, export limit and error recovery at 320/390/1440px')
finally:
    server.shutdown()
