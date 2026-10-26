"""Read-only, reproducible analysis of the supplied fastener workbook."""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
import hashlib
import json
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
SOURCE = ROOT / 'FTI_Smart_Fasteners_Price_List (2).xlsx'
SHEETS = ['SPL MS', 'SPL 8.8', 'SPL 10.9', 'SPL 12.9']


def number(value):
    return Decimal(str(value))


def analyze():
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
    products = {}
    bands = {}
    for title in SHEETS:
        entries = []
        for rownum, row in enumerate(workbook[title].iter_rows(values_only=True), 1):
            if rownum < 5 or not row[0] or not row[1] or not row[5]:
                continue
            match = re.fullmatch(r'\s*(\d+)\s*[Xx×]\s*(\d+)\s*', str(row[5]))
            assert match, (title, rownum, row[5])
            entries.append(dict(sheet=title, row=rownum, sku=str(row[0]),
                                family=row[1], grade=str(row[2]), finish=row[6],
                                diameter=int(match[1]), length=int(match[2]),
                                weight_kg_per_1000=row[3], quantity_kg=row[7], rate_per_kg=row[9]))
        assert len(entries) == 425
        assert len({r['sku'] for r in entries}) == 425
        assert len({(r['diameter'], r['length']) for r in entries}) == 425
        assert all(isinstance(r[k], (int, float)) and r[k] > 0
                   for r in entries for k in ['weight_kg_per_1000', 'rate_per_kg'])
        products[title] = entries
        groups = []
        for diameter in sorted({r['diameter'] for r in entries}):
            ordered = sorted((r for r in entries if r['diameter'] == diameter), key=lambda r: r['length'])
            current = None
            for row in ordered:
                if current is None or current['rate_per_kg'] != row['rate_per_kg']:
                    current = dict(diameter=diameter, min_length=row['length'],
                                   max_length=row['length'], rate_per_kg=row['rate_per_kg'],
                                   listed_lengths=[], source_rows=[])
                    groups.append(current)
                current['max_length'] = row['length']
                current['listed_lengths'].append(row['length'])
                current['source_rows'].append(row['row'])
        bands[title] = groups
        for row in entries:
            matches = [b for b in groups if b['diameter'] == row['diameter']
                       and b['min_length'] <= row['length'] <= b['max_length']]
            assert len(matches) == 1
            assert number(matches[0]['rate_per_kg']) == number(row['rate_per_kg'])

    base = products[SHEETS[0]]
    by_size = {(r['diameter'], r['length']): r for r in base}
    for title in SHEETS[1:]:
        for row in products[title]:
            original = by_size[(row['diameter'], row['length'])]
            assert number(original['weight_kg_per_1000']) == number(row['weight_kg_per_1000'])
    diameters = sorted({r['diameter'] for r in base})
    lengths = sorted({r['length'] for r in base})
    source_data = json.loads((ROOT / 'src/features/catalog/data/fastener-references.json').read_text(encoding='utf-8'))
    dimension_diameters = {int(r['size'].removeprefix('M')) for r in source_data['rows']
                          if r['chartId'] == 'bolt-dimensions'}
    metric = {(int(r['size'].removeprefix('M')), int(r['length'].replace(' mm', ''))): r
              for r in source_data['rows']
              if r['chartId'] == 'metric-weights' and r['groupId'] == 'hex-bolt'}
    matches, differences, no_match = [], [], []
    for row in base:
        ref = metric.get((row['diameter'], row['length']))
        if ref is None:
            no_match.append(row)
        elif number(ref['weightKgPer1000']) == number(row['weight_kg_per_1000']):
            matches.append(row)
        else:
            differences.append(dict(size=f"{row['diameter']} x {row['length']}", row=row['row'],
                                    workbook=row['weight_kg_per_1000'], reference=ref['weightKgPer1000'],
                                    reference_note=ref.get('reviewNote')))
    price_vectors = Counter(tuple(next(r['rate_per_kg'] for r in products[t]
                                      if (r['diameter'], r['length']) == key) for t in SHEETS)
                            for key in by_size)
    equal_8_10 = all(number(a['rate_per_kg']) == number(b['rate_per_kg'])
                    and (a['diameter'], a['length']) == (b['diameter'], b['length'])
                    for a, b in zip(products['SPL 8.8'], products['SPL 10.9']))
    result = dict(
        source=SOURCE.name, sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        source_bytes=SOURCE.stat().st_size,
        sheets=[dict(name=s.title, max_row=s.max_row, max_column=s.max_column) for s in workbook],
        product_rows=sum(map(len, products.values())), unique_sizes=len(by_size),
        diameters=diameters, lengths=lengths,
        unsupported_cross_product=len(diameters) * len(lengths) - len(by_size),
        product_rows_per_sheet={t:len(r) for t,r in products.items()},
        band_counts={t:len(b) for t,b in bands.items()},
        price_bands=bands, band_replay_matches=1700,
        unique_price_vectors=len(price_vectors), equal_8_8_and_10_9_rates=equal_8_10,
        shared_band_count=len(bands['SPL MS']) + len(bands['SPL 8.8']) + len(bands['SPL 12.9']),
        all_quantity_kg_50=all(r['quantity_kg']==50 for rs in products.values() for r in rs),
        missing_dimension_diameters=sorted(set(diameters)-dimension_diameters),
        dimension_covered_size_count=sum(r['diameter'] in dimension_diameters for r in base),
        weight_source_comparison=dict(equal=len(matches), different=len(differences),
                                      no_match=len(no_match), differences=differences),
        size_groups=[dict(diameter=d,count=sum(r['diameter']==d for r in base),
                          min_length=min(r['length'] for r in base if r['diameter']==d),
                          max_length=max(r['length'] for r in base if r['diameter']==d)) for d in diameters],
        source_rows=products,
        checks=['425 unique size pairs and SKUs per sheet', 'All primary prices and weights positive',
                'Same 425 size pairs and identical listed weights across all four sheets',
                'Exactly one compressed band matches each listed size within its sheet',
                'Every reconstructed rate equals the source rate using Decimal comparison']
    )
    workbook.close()
    (OUT / 'workbook-analysis.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k not in ['source_rows','price_bands','sheets']}, indent=2))
    return result


if __name__ == '__main__':
    analyze()
