"""Build a lossless table workbook from the saved, numbered web-reader extract."""
import json
import re
from pathlib import Path
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
source = json.loads((ROOT / 'source-tables.json').read_text(encoding='utf-8'))
names = [
    'Half bolts+nuts Metric 1', 'Half bolts+nuts Metric 2', 'Half bolts+nuts BSW',
    'Full bolts+nuts Metric 1', 'Full bolts+nuts Metric 2', 'Full bolts+nuts BSW',
    'Half bolts only Metric', 'Half bolts only BSW', 'Full bolts only Metric 1',
    'Full bolts only Metric 2', 'Full bolts only BSW', 'Round head rivets',
    'Square head set screws', 'Carriage bolts+nuts', 'Structural bolts',
    'Threaded rods', 'UNC studs', 'Nuts', 'Washers', 'Cheese head screws',
    'CSK screws IS', 'CSK screws BS', 'Pan Phillips screws', 'CSK Phillips screws',
    'Socket head cap screws', 'Socket set screws',
]

def clean(s):
    return s.replace('\u00a0', ' ').strip()

def is_table_line(s):
    return ' | ' in s and not s.startswith(('©', '【'))

tables, active = [], None
product = spec = unit = ''
for number, line in source['lines']:
    if line.startswith('#'):
        heading = line.lstrip('# ').strip()
        if heading.upper().startswith('APPROX'):
            unit = heading
        elif heading.startswith(('MM THREADS', 'BSW THREADS')):
            spec = heading
        else:
            product, spec, unit = heading, '', ''
    if is_table_line(line):
        if active is None:
            active = {'product': product, 'spec': spec, 'unit': unit, 'rows': []}
            tables.append(active)
        tokens = [clean(t) for t in line.split('|')]
        if not all(re.fullmatch(r':?-+:?', t) for t in tokens):
            active['rows'].append((number, tokens))
    else:
        active = None
assert len(tables) == len(names) == 26

wb = Workbook()
index = wb.active
index.title = 'Contents'
values = wb.create_sheet('All Values')
raw = wb.create_sheet('Source Cells')
values.append(['Chart', 'Product', 'Specification', 'Published unit', 'Row / size',
               'Column / type', 'Value', 'Published text', 'Source line'])
raw.append(['Chart', 'Source line', 'Source column', 'Published text'])
NAVY, BLUE, PALE, LINE = '17283B', '075299', 'EAF0FC', 'E5EAF0'
thin = Side(style='thin', color=LINE)
numeric = re.compile(r'^-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$')

def number_value(text):
    if not numeric.fullmatch(text):
        return text or None
    return float(text.replace(',', '')) if '.' in text else int(text.replace(',', ''))

def set_text(cell, value):
    cell.value = value or None
    if value:
        cell.data_type = 's'

def number_format(text):
    decimals = len(text.split('.')[1]) if '.' in text else 0
    return ('#,##0' if ',' in text else '0') + ('.' + '0' * decimals if decimals else '')

def decorate(ws, header_row, widths, freeze):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    for c in ws[header_row]:
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

index.append(['ITA Fasteners — weight charts'])
index.append(['Source', source['url']])
index['B2'].hyperlink = source['url']
index.append(['Extracted (UTC)', source['extracted_at']])
index.append(['Coverage', 'All 26 weight tables on the linked page. Each chart has its own sheet.'])
index.append(['Units', 'Keep the unit shown on each sheet: kg/100 pieces, kg/piece, or kg/metre/piece.'])
index.append(['Blank values', 'Blank means no published value; it is not zero.'])
index.append(['Numbers', 'Weights are numeric. Size labels and fractions are text. Source Cells retains every original table cell.'])
index.append(['Scope', 'Table headings, standards, units and table values; navigation, marketing prose and decorative images excluded.'])
index.append(['Preservation', 'Published values and axis labels are retained, including any source inconsistencies; no interpolation or correction.'])
index.append([])
index.append(['Sheet', 'Product', 'Specification', 'Published unit', 'Data rows', 'Value cells', 'Source lines'])
index_checks, source_checks = [], []
total_rows = total_values = 0

for i, (table, name) in enumerate(zip(tables, names), 1):
    ws = wb.create_sheet(f'{i:02d} {name}')
    original = table['rows']
    # Retain an exact copy of each source token, including blank cells.
    for ln, cells in original:
        for col, value in enumerate(cells, 1):
            raw.append([ws.title, ln, col, None])
            set_text(raw.cell(raw.max_row, 4), value)
            source_checks.append(value or None)
    rows = [(ln, list(cells)) for ln, cells in original]
    if rows[0][1] == ['LENGTH', 'DIAMETER']:
        # Web reader elides the spanning empty cell before diameter headers.
        rows[1][1].insert(0, '')
        header_count = 2
    elif name == 'Nuts':
        header_count = 3
    elif name == 'Washers':
        header_count = 5
    else:
        header_count = 1
    width = max(len(r) for _, r in rows)
    for ln, cells in rows:
        cells.extend([''] * (width - len(cells)))
    for r, text in [(1, table['product']), (2, table['spec']), (3, table['unit'])]:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        set_text(ws.cell(r, 1), text)
        ws.cell(r, 1).font = Font(name='Calibri', size=15 if r == 1 else 11,
                                  bold=r in (1, 3), color=NAVY if r == 1 else BLUE)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[r].height = 40 if len(text) > 70 else 25
    ws.cell(4, 1, 'Source page')
    ws.cell(4, 2, source['url']).hyperlink = source['url']
    ws.cell(4, 2).font = Font(color=BLUE, underline='single')
    top = 6
    secondary_size = 8 if name == 'Nuts' else 10 if name == 'Washers' else 2 if name == 'Threaded rods' else None
    size_columns = {0, secondary_size} if secondary_size is not None else {0}
    value_count = 0
    for ri, (ln, cells) in enumerate(rows):
        excel_row = top + ri
        is_header = ri < header_count
        for ci, text in enumerate(cells):
            cell = ws.cell(excel_row, ci + 1)
            if is_header or ci in size_columns:
                set_text(cell, text)
            else:
                cell.value = number_value(text)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = number_format(text)
                elif cell.value:
                    cell.data_type = 's'
            cell.border = Border(bottom=thin, right=thin)
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left' if ci in size_columns else 'right')
            cell.font = Font(name='Calibri', size=11, color='FFFFFF' if is_header else NAVY, bold=is_header or ci == 0)
            if is_header:
                cell.fill = PatternFill('solid', fgColor=NAVY)
            elif ri % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F5F7FB')
            if not is_header and ci not in size_columns and text:
                row_label = cells[secondary_size] if secondary_size is not None and ci > secondary_size else cells[0]
                labels = [rows[h][1][ci] for h in range(header_count) if rows[h][1][ci]]
                # Grouped washer headings are made explicit only in All Values.
                if name == 'Washers' and 4 <= ci <= 7:
                    labels.insert(0, 'SQUARE TAPER WASHERS')
                column = ' / '.join(labels)
                values.append([ws.title, table['product'], table['spec'], table['unit'], row_label, column,
                               number_value(text), text, ln])
                values.cell(values.max_row, 7).number_format = number_format(text) if numeric.fullmatch(text) else 'General'
                set_text(values.cell(values.max_row, 8), text)
                value_count += 1
        ws.row_dimensions[excel_row].height = 48 if is_header else 23
    if rows[0][1][:2] == ['LENGTH', 'DIAMETER']:
        ws.merge_cells(start_row=top, start_column=2, end_row=top, end_column=width)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f'B{top + header_count}'
    for ci in range(1, width + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 22 if ci == 1 else 19 if name in ('Nuts', 'Washers') else 13
    ws.print_title_rows = f'1:{top + header_count - 1}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    data_rows = len(rows) - header_count
    total_rows += data_rows
    total_values += value_count
    index.append([ws.title, table['product'], table['spec'], table['unit'], data_rows, value_count,
                  f'{original[0][0]}–{original[-1][0]}'])
    index.cell(index.max_row, 1).hyperlink = f"#'{ws.title}'!A1"
    index.cell(index.max_row, 1).font = Font(color=BLUE, underline='single')
    index_checks.append((ws.title, data_rows, value_count))

for ws, widths, table_name in [(values, [32, 55, 38, 45, 22, 65, 16, 18, 14], 'WeightValues'),
                              (raw, [32, 15, 16, 65], 'OriginalCells')]:
    decorate(ws, 1, widths, 'A2')
    ws.auto_filter.ref = ws.dimensions
    t = Table(displayName=table_name, ref=ws.dimensions)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
decorate(index, 11, [34, 65, 38, 48, 14, 14, 18], 'A12')
index.merge_cells('A1:G1')
index['A1'].font = Font(size=20, bold=True, color=NAVY)
index.row_dimensions[1].height = 34
for row in index.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)
for r in range(4, 10):
    index.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    index.row_dimensions[r].height = 28
index.auto_filter.ref = f'A11:G{index.max_row}'
index.append(['TOTAL', '', '', '', total_rows, total_values])
# Keep the chart sheets immediately after Contents.
wb._sheets = [index] + [wb[n] for n, _, _ in index_checks] + [values, raw]
output = ROOT / 'ITA_Fasteners_Weight_Charts.xlsx'
wb.save(output)

# Reopen the workbook and verify that no table tokens were lost during export.
check = load_workbook(output, data_only=False)
assert len(check.sheetnames) == 29
assert check['All Values'].max_row - 1 == total_values
actual = [row[0] for row in check['Source Cells'].iter_rows(min_row=2, min_col=4, max_col=4, values_only=True)]
assert actual == source_checks
assert not any(c.data_type == 'f' for ws in check for row in ws for c in row)
report = {'workbook': str(output), 'charts': 26, 'worksheets': 29, 'data_rows': total_rows,
          'value_cells': total_values, 'source_cells': len(source_checks), 'source': source['url'],
          'verified': 'Workbook reopened; every original table token retained; no formulas introduced.'}
(ROOT / 'extraction-report.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
print(json.dumps(report, indent=2))
